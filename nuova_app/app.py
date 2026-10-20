import json
import io
from pathlib import Path
import streamlit as st
import pandas as pd
import fitz
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent
PDF_PATH = ROOT / "assets" / "Riepilogo chiusura strade.pdf"
with open(ROOT / "data" / "tavole.json", encoding="utf-8") as f:
    TAVOLE = json.load(f)
with open(ROOT / "data" / "segnali.json", encoding="utf-8") as f:
    SEGNALI = json.load(f)

st.set_page_config(page_title="Tavole e abaco segnaletica", page_icon="🚧", layout="wide")
st.title("Catalogo tavole e abaco segnaletica")
st.caption("Prima versione operativa: catalogo PDF + archivio segnali Excel + abaco modificabile ed esportabile. Le quantità non vengono dedotte automaticamente dagli schemi: vanno verificate/inserite dall’utente.")

with st.sidebar:
    st.header("Dati di cantiere")
    commessa = st.text_input("Commessa / cantiere")
    committente = st.text_input("Committente")
    strada = st.text_input("Strada / località")
    progressiva_da = st.text_input("Progressiva iniziale")
    progressiva_a = st.text_input("Progressiva finale")
    direzione = st.text_input("Direzione / carreggiata")
    note = st.text_area("Note", height=90)

tab1, tab2, tab3 = st.tabs(["Catalogo tavole", "Componi abaco", "Archivio segnali"])

with tab1:
    st.subheader("Tavole del PDF originale")
    df_t = pd.DataFrame([{"Pagina PDF":x["pagina_pdf"], "Tavola":x["tavola"], "Titolo estratto":x["titolo_estratto"]} for x in TAVOLE])
    c1,c2 = st.columns([1,2])
    with c1:
        filtro = st.text_input("Cerca numero o titolo")
    with c2:
        solo_tavole = st.checkbox("Mostra solo pagine con intestazione TAVOLA", value=True)
    shown=df_t.copy()
    if solo_tavole:
        shown=shown[shown["Tavola"].astype(str).str.len()>0]
    if filtro:
        shown=shown[shown.astype(str).apply(lambda col: col.str.contains(filtro,case=False,na=False)).any(axis=1)]
    st.dataframe(shown, use_container_width=True, hide_index=True, height=320)
    pages_options=[x["pagina_pdf"] for x in TAVOLE]
    page_num=st.selectbox("Apri pagina PDF", pages_options, index=0, format_func=lambda n: f"Pagina {n}"+ (f" — Tavola {next((x['tavola'] for x in TAVOLE if x['pagina_pdf']==n and x['tavola']), '')}" if any(x['pagina_pdf']==n and x['tavola'] for x in TAVOLE) else ""))
    doc=fitz.open(PDF_PATH)
    page=doc[page_num-1]
    pix=page.get_pixmap(matrix=fitz.Matrix(1.35,1.35), alpha=False)
    st.image(Image.open(io.BytesIO(pix.tobytes("png"))), caption=f"Pagina {page_num} del PDF originale", use_container_width=True)
    with st.expander("Testo estratto dalla pagina"):
        st.text(page.get_text("text") or "Nessun testo estraibile: la pagina potrebbe essere composta da immagini.")

with tab2:
    st.subheader("Abaco personalizzabile")
    st.info("Seleziona i segnali dall’archivio e inserisci quantità/dettagli. L’associazione tra tavola e segnali non è ancora certificata automaticamente; controllare lo schema PDF prima dell’esportazione.")
    choices=[]
    for s in SEGNALI:
        label=" — ".join(str(v) for v in [s.get("figura"),s.get("descrizione")] if v not in (None,""))
        if not label:
            label=f"Riga Excel {s['riga_excel']}"
        choices.append((s["id"],label))
    labels={label:ident for ident,label in choices}
    selected_labels=st.multiselect("Segnali da inserire nell’abaco", [x[1] for x in choices])
    selected_ids=[labels[x] for x in selected_labels]
    byid={s["id"]:s for s in SEGNALI}
    rows=[]
    for ident in selected_ids:
        s=byid[ident]
        rows.append({"Codice":s.get("codice_originale") or "", "Figura":s.get("figura") or "", "Descrizione":s.get("descrizione") or "", "Categoria":s.get("categoria") or "", "Dettagli":s.get("dettagli") or "", "Quantità":1, "Note":""})
    if rows:
        edited=st.data_editor(pd.DataFrame(rows), use_container_width=True, num_rows="dynamic", key="abaco_editor")
    else:
        edited=pd.DataFrame(columns=["Codice","Figura","Descrizione","Categoria","Dettagli","Quantità","Note"])
        st.caption("Seleziona almeno un segnale per iniziare.")
    if st.button("Genera file Excel", type="primary", disabled=edited.empty):
        wb=Workbook()
        ws=wb.active
        ws.title="Abaco"
        ws.merge_cells("A1:G1")
        ws["A1"]="ABACO SEGNALETICA DI CANTIERE"
        ws["A1"].font=Font(bold=True,size=16,color="FFFFFF")
        ws["A1"].fill=PatternFill("solid",fgColor="17365D")
        ws["A1"].alignment=Alignment(horizontal="center")
        meta=[("Commessa",commessa),("Committente",committente),("Strada / località",strada),("Progressiva iniziale",progressiva_da),("Progressiva finale",progressiva_a),("Direzione / carreggiata",direzione),("Tavola di riferimento",str(next((x['tavola'] for x in TAVOLE if x['pagina_pdf']==page_num and x['tavola']),''))),("Note",note)]
        for i,(k,v) in enumerate(meta,3):
            ws.cell(i,1,k).font=Font(bold=True)
            ws.cell(i,2,v)
        start=3+len(meta)+1
        headers=list(edited.columns)
        for j,h in enumerate(headers,1):
            c=ws.cell(start,j,h); c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="4472C4")
        for i,row in enumerate(edited.itertuples(index=False,name=None),start+1):
            for j,val in enumerate(row,1):
                ws.cell(i,j,val)
        for col,width in enumerate([18,16,42,25,24,14,32],1):
            ws.column_dimensions[get_column_letter(col)].width=width
        ws.freeze_panes=f"A{start+1}"
        ws.auto_filter.ref=f"A{start}:{get_column_letter(len(headers))}{start+len(edited)}"
        for row in ws.iter_rows(min_row=start+1,max_row=start+len(edited),min_col=1,max_col=len(headers)):
            for cell in row: cell.alignment=Alignment(vertical="top",wrap_text=True)
        bio=io.BytesIO(); wb.save(bio); bio.seek(0)
        st.download_button("Scarica abaco Excel",bio.getvalue(),file_name="abaco_segnaletica.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

with tab3:
    st.subheader("Archivio dei segnali disponibili")
    q=st.text_input("Filtra per figura, descrizione o categoria", key="q_signali")
    catalog=[]
    for s in SEGNALI:
        catalog.append({"Figura":s.get("figura") or "", "Descrizione":s.get("descrizione") or "", "Categoria":s.get("categoria") or "", "Dettagli":s.get("dettagli") or "", "Riga Excel":s.get("riga_excel"), "ID":s.get("id")})
    d=pd.DataFrame(catalog)
    if q:
        d=d[d.astype(str).apply(lambda col: col.str.contains(q,case=False,na=False)).any(axis=1)]
    st.dataframe(d,use_container_width=True,hide_index=True,height=420)
    st.caption(f"Archivio importato: {len(SEGNALI)} voci. Il file Excel originale è conservato in assets per confronto.")
