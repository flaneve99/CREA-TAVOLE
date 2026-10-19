import io
import os
import urllib.request
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageDraw, ImageFont
import streamlit as st

# Importiamo il dizionario di configurazione che hai appena creato
from config_tavole import CONFIG_TAVOLE

# ==================================================
# CONFIGURAZIONE STREAMLIT E PERCORSI
# ==================================================
st.set_page_config(
    page_title="Gestione Cantieri Stradali v2 - Automatico",
    layout="wide",
)

cartella_app = os.path.dirname(os.path.abspath(__file__))
dir_immagini = os.path.join(cartella_app, "immagini_segnali")

st.title("🚧 Gestione Cantieri Stradali (Motore Automatico v2)")
st.write("Seleziona la Tavola: l'interfaccia, l'Abaco Excel e la Tavola JPG si adatteranno automaticamente.")

# ==================================================
# CARICAMENTO FONT COMPATIBILE
# ==================================================
@st.cache_resource
def load_custom_font(size):
    possibili_font = [
        os.path.join(cartella_app, "DejaVuSans-Bold.ttf"),
        "DejaVuSans-Bold.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf",
        "C:/Windows/Fonts/arialbd.ttf",
        "/Library/Fonts/Arial Bold.ttf",
    ]
    for percorso in possibili_font:
        if os.path.isfile(percorso):
            try:
                return ImageFont.truetype(percorso, size)
            except Exception:
                continue
    try:
        return ImageFont.load_default(size=size)
    except TypeError:
        return ImageFont.load_default()

# ==================================================
# FUNZIONE UNIVERSALE PER LA GENERAZIONE DELL'EXCEL
# ==================================================
def genera_abaco_universale(chiave_tavola, dati_segnali):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Abaco Segnaletica"
    
    ws.merge_cells("A1:E1")
    titolo_principale = ws["A1"]
    titolo_principale.value = f"TABELLA RIEPILOGATIVA SEGNALETICA DI CANTIERE - {chiave_tavola.split(' - ')[0]}"
    
    fill_verde_chiaro = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
    fill_verde_scuro = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
    font_titolo = Font(name="Arial", size=11, bold=True, color="000000")
    font_header = Font(name="Arial", size=10, bold=True, color="000000")
    font_testo = Font(name="Arial", size=10, color="000000")
    bordo_sottile = Border(
        left=Side(style="thin", color="000000"), right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"), bottom=Side(style="thin", color="000000")
    )
    
    titolo_principale.fill = fill_verde_chiaro
    titolo_principale.font = font_titolo
    titolo_principale.alignment = Alignment(horizontal="center", vertical="center")
    
    for col in range(1, 6):
        ws.cell(row=1, column=col).border = bordo_sottile
        
    headers = ["SEGNALE", "FIGURA", "DESCRIZIONE", "DETTAGLI", "SEGNALETICA"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx)
        cell.value = header
        cell.fill = fill_verde_scuro
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = bordo_sottile
        
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 45
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 25
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 25
    
    immagini_inserite = 0
    for i, row_data in enumerate(dati_segnali, start=3):
        img_nome, figura, descrizione, dettagli, segnaletica = row_data
        ws.row_dimensions[i].height = 55
        c_fig = ws.cell(row=i, column=2, value=figura)
        c_desc = ws.cell(row=i, column=3, value=descrizione)
        c_det = ws.cell(row=i, column=4, value=dettagli)
        c_seg = ws.cell(row=i, column=5, value=segnaletica)
        
        for cell in [c_fig, c_desc, c_det, c_seg]:
            cell.font = font_testo
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = bordo_sottile
            
        c_segnale = ws.cell(row=i, column=1)
        c_segnale.border = bordo_sottile
        
        nomi_da_cercare = [img_nome, img_nome.replace(".jpg", ".png"), img_nome.replace(".png", ".jpg")]
        percorso_trovato = None
        for nome in nomi_da_cercare:
            tentativo = os.path.join(dir_immagini, nome)
            if os.path.exists(tentativo):
                percorso_trovato = tentativo
                break
                
        if percorso_trovato:
            try:
                img_pil = Image.open(percorso_trovato).convert("RGB")
                temp_path = os.path.join("/tmp", f"abaco_auto_img_{i}.jpg")
                img_pil.save(temp_path, "JPEG")
                img_obj = OpenpyxlImage(temp_path)
                marker_from = AnchorMarker(col=0, colOff=150000, row=i-1, rowOff=40000)
                marker_to = AnchorMarker(col=1, colOff=-150000, row=i, rowOff=-40000)
                img_obj.anchor = TwoCellAnchor(editAs="twoCell", _from=marker_from, to=marker_to)
                ws.add_image(img_obj)
                immagini_inserite += 1
            except Exception as e:
                st.warning(f"Errore caricamento immagine per {figura}: {e}")
        else:
            c_segnale.value = "[Img non trovata]"
            c_segnale.font = Font(name="Arial", size=8, color="FF0000")
            c_segnale.alignment = Alignment(horizontal="center", vertical="center")
            
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue(), immagini_inserite

# ==================================================
# FUNZIONE UNIVERSALE PER COMPILARE IL JPG
# ==================================================
def compila_immagine_universale(config_tavola, valori_km, testo_ubicazione):
    img_path = os.path.join(cartella_app, config_tavola["immagine_sfondo"])
    if not os.path.exists(img_path):
        st.error(f"Immagine sfondo '{config_tavola['immagine_sfondo']}' non trovata.")
        return None
        
    with Image.open(img_path) as img_orig:
        base_img = img_orig.convert("RGB")
        
    width, height = base_img.size
    draw = ImageDraw.Draw(base_img)
    
    font_km = load_custom_font(config_tavola["moltiplicatore_font_km"])
    font_ub = load_custom_font(config_tavola["moltiplicatore_font_ub"])
    BORDO = 2
    
    # 1. Scrittura automatica delle progressive chilometriche inserite
    for inp in config_tavola["inputs_km"]:
        testo = f"KM {valori_km[inp['id']]}"
        x = int(width * inp["x_pct"])
        y = int(height * inp["y_pct"])
        draw.text((x, y), testo, fill=(255, 0, 0), font=font_km, stroke_width=BORDO, stroke_fill=(255, 255, 255))
        
    # 2. Scrittura automatica dell'ubicazione multilinea
    pos_ub = config_tavola["posizione_ubicazione"]
    x_ub = int(width * pos_ub["x_pct"])
    y_ub = int(height * pos_ub["y_pct"])
    draw.multiline_text((x_ub, y_ub), testo_ubicazione, fill=(255, 0, 0), font=font_ub, align="center", spacing=6, stroke_width=BORDO, stroke_fill=(255, 255, 255))
    
    img_io = io.BytesIO()
    base_img.save(img_io, format="JPEG", quality=100, subsampling=0)
    return img_io.getvalue(), base_img

# ==================================================
# INTERFACCIA STREAMLIT DINAMICA (DATA-DRIVEN)
# ==================================================
tavola_selezionata = st.selectbox("Seleziona la Tavola di Cantiere", list(CONFIG_TAVOLE.keys()))
tavola = CONFIG_TAVOLE[tavola_selezionata]

st.subheader(f"📍 Configurazione: {tavola_selezionata.split(' - ')[0]}")
ubicazione_testo = st.text_area("Ubicazione / Direzione (Può essere divisa su più righe)", value=tavola["ubicazione_default"])

# Generazione dinamica dei campi di input KM su due colonne
st.write("### 🛣️ Progressive Chilometriche")
valori_km = {}
col1, col2 = st.columns(2)

for idx, campo in enumerate(tavola["inputs_km"]):
    with col1 if idx % 2 == 0 else col2:
        valori_km[campo["id"]] = st.text_input(campo["label"], value=campo["default"])

# Pulsante per Excel Universale
if st.button(f"Genera Abaco Segnaletica (Excel) - {tavola_selezionata.split(' - ')[0]}"):
    excel_bytes, n_img = genera_abaco_universale(tavola_selezionata, tavola["abaco_segnali"])
    st.success(f"Abaco Excel generato con successo! (Immagini inserite: {n_img}/{len(tavola['abaco_segnali'])})")
    st.download_button(
        label="📥 Scarica Excel Abaco Segnaletica Formattato",
        data=excel_bytes,
        file_name=f"Abaco_Segnaletica_{tavola_selezionata.split(' - ')[0].replace(' ', '_')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

st.divider()

# Compilazione ed anteprima automatica del JPG
st.subheader("🖼️ Anteprima Tavola Personalizzata")
risultato_jpg = compila_immagine_universale(tavola, valori_km, ubicazione_testo)

if risultato_jpg is not None:
    jpg_bytes, img_mostrare = risultato_jpg
    st.image(img_mostrare, caption=f"{tavola_selezionata} - Configurazione completata", use_container_width=True)
    st.download_button(
        label="📥 Scarica Tavola in formato JPG",
        data=jpg_bytes,
        file_name=f"{tavola_selezionata.split(' - ')[0].replace(' ', '_')}_Compilata.jpg",
        mime="image/jpeg"
    )
