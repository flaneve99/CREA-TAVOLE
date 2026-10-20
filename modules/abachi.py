import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

def crea_excel_abaco(titolo_abaco, headers, righe_dati):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Abaco"

    # Stili
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    font_data = Font(name="Calibri", size=10)
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # Intestazione Titolo
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    cell_title = ws.cell(row=1, column=1, value=titolo_abaco)
    cell_title.font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    cell_title.alignment = Alignment(horizontal='left', vertical='center')

    # Intestazioni Colonne
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num, value=header)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center

    # Dati
    for row_idx, riga in enumerate(righe_dati, start=4):
        for col_idx, val in enumerate(riga, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_data
            cell.border = border_thin
            cell.alignment = align_center

    # Autofit colonne
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
